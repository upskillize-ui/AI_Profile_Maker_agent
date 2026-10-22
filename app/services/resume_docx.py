"""
Resume export (DOCX) — v13
═══════════════════════════
Builds a plain, ATS-friendly Word resume from the same context the profile
page renders from: white background, black text, one accent-free column.
No images, no colour, no tables — recruiters' parsers read it cleanly.

Only sections with data are written (same rule as the profile page).
"""

from io import BytesIO
from typing import Any, Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor

_BLACK = RGBColor(0, 0, 0)
_FONT = "Calibri"


def _fmt_date(s) -> str:
    if not s:
        return ""
    d = str(s).split("T")[0]
    if len(d) >= 10 and d[4] == "-":
        months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
        try:
            return f"{int(d[8:10])} {months[int(d[5:7]) - 1]} {d[0:4]}"
        except (ValueError, IndexError):
            return d
    return d


def _run(p, text, bold=False, size=11, italic=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = _FONT
    r.font.color.rgb = _BLACK
    return r


def _para(doc, text="", bold=False, size=11, align=None, space_after=4, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    if text:
        _run(p, text, bold=bold, size=size, italic=italic)
    return p


def _heading(doc, text):
    p = _para(doc, text.upper(), bold=True, size=11.5, space_after=2)
    p.paragraph_format.space_before = Pt(10)
    # thin rule under the heading
    pPr = p._p.get_or_add_pPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def _bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    _run(p, text, size=size)
    return p


def _line(doc, left: str, right: str = "", bold_left=True, size=11):
    """Two-part line: left bold, right plain (dates / scores)."""
    p = _para(doc, space_after=1)
    _run(p, left, bold=bold_left, size=size)
    if right:
        _run(p, f"   {right}", size=size - 0.5)
    return p


def _skills_text(skills: Dict) -> List[str]:
    out = []
    for grp, label in (("technical_skills", "Technical"), ("soft_skills", "Soft skills"), ("tools", "Tools")):
        items = []
        seen = set()
        for sk in (skills.get(grp) or []):
            nm = (sk.get("name") if isinstance(sk, dict) else sk) or ""
            nm = str(nm).strip()
            if len(nm) > 1 and nm.lower() not in seen:
                seen.add(nm.lower())
                items.append(nm)
        if items:
            out.append(f"{label}: " + ", ".join(items))
    return out


def build_resume_docx(ctx: Dict[str, Any]) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    style = doc.styles["Normal"]
    style.font.name = _FONT
    style.font.size = Pt(11)
    style.font.color.rgb = _BLACK

    # ── Header ──
    _para(doc, ctx.get("student_name") or "", bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    if ctx.get("headline"):
        _para(doc, ctx["headline"], size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    contact = [x for x in (ctx.get("student_location"), ctx.get("student_email"), ctx.get("student_phone")) if x]
    links = [x for x in (ctx.get("linkedin_url"), ctx.get("github_url"), ctx.get("portfolio_url")) if x]
    if contact:
        _para(doc, "  |  ".join(contact), size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    if links:
        _para(doc, "  |  ".join(links), size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    # ── Summary ──
    summary = (ctx.get("professional_summary") or "").strip()
    if len(summary) > 20:
        _heading(doc, "Professional Summary")
        lines = [ln.strip() for ln in summary.replace("•", "\n").split("\n") if ln.strip()]
        if len(lines) > 1:
            for ln in lines:
                _bullet(doc, ln)
        else:
            _para(doc, summary, size=10.5)

    # ── Skills ──
    skills = ctx.get("skills_data") or {}
    sk_lines = _skills_text(skills)
    if sk_lines:
        _heading(doc, "Skills")
        for ln in sk_lines:
            _para(doc, ln, size=10.5, space_after=2)

    # ── Experience ──
    work = ctx.get("work_experience") or []
    if work:
        _heading(doc, "Experience")
        for exp in work:
            _line(doc, f"{exp.get('title') or ''}{(' — ' + exp['company']) if exp.get('company') else ''}", exp.get("duration") or "")
            if exp.get("description"):
                _bullet(doc, exp["description"])

    # ── Projects ──
    projects = ctx.get("projects_data") or []
    if projects:
        _heading(doc, "Projects")
        for pr in projects[:6]:
            techs = pr.get("technologies_used")
            if isinstance(techs, (list, tuple)):
                techs = ", ".join(str(t).strip() for t in techs if str(t).strip())
            _line(doc, pr.get("name") or pr.get("title") or "", f"({techs})" if techs else "")
            if pr.get("description"):
                _bullet(doc, pr["description"])

    # ── Education ──
    edu = ctx.get("education_data") or []
    if edu:
        _heading(doc, "Education")
        for e in edu:
            deg = e.get("degree") or ""
            fos = e.get("field_of_study") or e.get("specialization")
            if fos:
                deg = f"{deg} – {fos}"
            right = " · ".join(x for x in (
                e.get("institution"),
                str(e.get("year")) if e.get("year") else "",
                (f"{e['percentage']}%" if e.get("percentage") and "%" not in str(e["percentage"]) else (str(e["percentage"]) if e.get("percentage") else "")) or (f"{e['cgpa']} CGPA" if e.get("cgpa") else ""),
            ) if x)
            _line(doc, deg, right)

    # ── Certifications ──
    certs = ctx.get("certifications_data") or []
    certs = [c for c in certs if (c.get("certificate_name") or c.get("name"))]
    if certs:
        _heading(doc, "Certifications")
        for c in certs:
            right = " · ".join(x for x in (c.get("issuer"), _fmt_date(c.get("issued_at")),
                                           "Upskillize-verified" if c.get("upskillize_verified") else "") if x)
            _line(doc, c.get("certificate_name") or c.get("name"), right)

    # ── Validated performance (Upskillize LMS) ──
    perf_rows: List[str] = []
    for key, label in (("assessments", "Assessment"), ("mock_tests", "Mock test"),
                       ("assignments", "Assignment"), ("case_studies", "Case study"),
                       ("capstones", "Capstone"), ("mock_interviews", "Mock interview")):
        src = {"assessments": "test_scores_raw", "mock_tests": "mock_tests_raw", "assignments": "assignments_raw",
               "case_studies": "case_studies_raw", "capstones": "capstone_projects",
               "mock_interviews": "mock_interviews_raw"}[key]
        for r in (ctx.get(src) or [])[:3]:
            sc = r.get("score_pct")
            title = r.get("title") or r.get("project_name") or r.get("role") or ""
            if title and sc is not None:
                perf_rows.append(f"{label}: {title} — {int(round(float(sc)))}/100")
    for s in (ctx.get("industry_sessions_raw") or [])[:3]:
        who = " · ".join(x for x in (s.get("speaker"), s.get("company")) if x)
        perf_rows.append(f"Industry session: {s.get('title') or ''}{(' (' + who + ')') if who else ''}{(' — ' + _fmt_date(s.get('held_at'))) if s.get('held_at') else ''}")
    for h in (ctx.get("hackathons_raw") or [])[:2]:
        perf_rows.append(f"Hackathon: {h.get('title') or h.get('name') or ''}{(' — ' + (h.get('badge') or '')) if h.get('badge') else ''}")
    punct = ((ctx.get("perf_snapshot") or {}).get("punctuality") or {})
    if punct.get("score") is not None:
        perf_rows.append(f"Punctuality score: {int(round(float(punct['score'])))}/100{(' (' + punct['band'] + ')') if punct.get('band') else ''}")
    if perf_rows:
        _heading(doc, "Validated Performance — Upskillize LMS")
        for ln in perf_rows:
            _bullet(doc, ln)

    # ── Courses ──
    courses = ctx.get("courses_data") or []
    done = [c for c in courses if str(c.get("completion_status") or "").lower() == "completed"
            or float(c.get("progress_percentage") or 0) >= 100]
    if done:
        _heading(doc, "Completed Programmes")
        for c in done:
            _line(doc, c.get("course_name") or "", f"Upskillize{(' · ' + _fmt_date(c.get('completed_at'))) if c.get('completed_at') else ''}")

    # ── Open to / goals ──
    jp = ctx.get("job_preferences") or {}
    goal = ctx.get("career_goal_line") or ctx.get("personal_career_goals")
    open_bits = [x for x in (jp.get("preferred_role"), jp.get("preferred_location"),
                             jp.get("work_mode"), jp.get("employment_type"),
                             (f"Notice: {jp['notice_period']}" if jp.get("notice_period") else "")) if x]
    if open_bits or goal:
        _heading(doc, "Career")
        if open_bits:
            _para(doc, "Open to: " + " · ".join(open_bits), size=10.5, space_after=2)
        if goal:
            _para(doc, goal, size=10.5, space_after=2)

    # ── Footer note ──
    p = _para(doc, "", space_after=0)
    p.paragraph_format.space_before = Pt(14)
    _run(p, f"Generated from the Upskillize AI Profile · {ctx.get('generated_date', '')}".title().replace("Ai", "AI"),
         size=8.5, italic=True)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
